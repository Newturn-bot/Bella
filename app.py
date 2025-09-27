# app.py
import os
import time
import streamlit as st
from dotenv import load_dotenv
import google.generativeai as genai
import PyPDF2

# Initialize Streamlit session state for messages
if "messages" not in st.session_state:
    st.session_state["messages"] = []

def create_word(text):
    doc = Document()
    doc.add_heading("CV Summary", level=1)
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def create_pdf(text):
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    text_obj = c.beginText(40, 750)
    text_obj.setFont("Helvetica", 12)

    for line in text.split("\n"):
        text_obj.textLine(line)
    c.drawText(text_obj)
    c.save()
    buffer.seek(0)
    return buffer
# ----------------------------
# 1. Set page config (MUST be first Streamlit command)
# -----------------------------
st.set_page_config(
    page_title="Newturn Chatbot",
    page_icon="🤖",
    layout="wide"  # or "centered"
)

# -----------------------------
# 2. Load API key
# -----------------------------
load_dotenv()
api_key = st.secrets["GOOGLE_API_KEY"]

if not api_key:
    raise ValueError("❌ GOOGLE_API_KEY not found in Streamlit secrets")

genai.configure(api_key=api_key)

# -----------------------------
# 3. Streamlit UI
# -----------------------------
st.title("🤖 Newturn Chatbot")
st.write("Hello 👋 I'm Newturn, your smart CV assistant. Upload your CV and ask me anything!")

# -----------------------------
# 4. PDF Upload & Text Extraction
# -----------------------------
uploaded_file = st.file_uploader("📄 Upload  CV (PDF only)", type=["pdf"])

cv_text = ""
if uploaded_file:
    pdf_reader = PyPDF2.PdfReader(uploaded_file)
    for page in pdf_reader.pages:
        cv_text += page.extract_text()
    st.success("✅ CV uploaded and processed!")
# -----------------------------
# 5. Summarize CV Button
# -----------------------------
from io import BytesIO
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

if uploaded_file:
    if st.button("📊 Summarize this CV"):
        try:
            model = genai.GenerativeModel("models/gemini-2.0-flash-001")

            summary_prompt = f"""
            You are Newturn, a helpful CV analysis assistant.

            Please summarize this CV in a structured format:
            - Candidate Name (if available)
            - Education
            - Work Experience
            - Technical Skills
            - Languages
            - Certifications
            - Soft Skills

            CV Content:
            {cv_text}
            """

            response = model.generate_content(summary_prompt)
            summary_text = response.text

            st.subheader("📑 CV Summary")
            st.write(summary_text)

            # -----------------------------
            # Export Options
            # -----------------------------
            col1, col2 = st.columns(2)

            with col1:
                if st.download_button("⬇️ Download as Word (docx)", 
                                      data=create_word(summary_text),
                                      file_name="cv_summary.docx"):
                    st.success("✅ Word file downloaded!")

            with col2:
                if st.download_button("⬇️ Download as PDF", 
                                      data=create_pdf(summary_text),
                                      file_name="cv_summary.pdf"):
                    st.success("✅ PDF file downloaded!")

        except Exception as e:
            st.error(f"⚠️ Error during summarization: {e}")
# -----------------------------
# 6. Chat Input
# -----------------------------
if user_input := st.chat_input("💬 Ask me anything about the CV:"):
    # Save user input
    st.session_state.messages.append({"role": "user", "content": user_input})
    st.chat_message("user").write(user_input)

    try:
        model = genai.GenerativeModel("models/gemini-2.0-flash-001")

        # Build conversation context (CV + history)
        context = cv_text if cv_text else ""
        history = "\n".join([f"{m['role']}: {m['content']}" for m in st.session_state.messages])

        prompt = f"""
        Here is the CV:
        {context}

        Conversation so far:
        {history}

        Assistant, continue the conversation.
        """

        response = model.generate_content(prompt)

        reply = response.text
        st.session_state.messages.append({"role": "assistant", "content": reply})
        st.chat_message("assistant").write(reply)

    except Exception as e:    
        st.error(f"⚠️ Error: {e}")

